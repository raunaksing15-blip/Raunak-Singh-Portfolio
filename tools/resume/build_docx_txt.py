from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from _profile import PROFILE as P, HERE, ASSETS
OUT=str(ASSETS)+"/"
def hyper(par,url,text):
    r_id=par.part.relate_to(url,"http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",is_external=True)
    h=OxmlElement("w:hyperlink"); h.set(qn("r:id"),r_id)
    r=OxmlElement("w:r"); rp=OxmlElement("w:rPr")
    c=OxmlElement("w:color"); c.set(qn("w:val"),"1A4FB4"); u=OxmlElement("w:u"); u.set(qn("w:val"),"single")
    rp.append(c); rp.append(u); r.append(rp); t=OxmlElement("w:t"); t.text=text; r.append(t); h.append(r); par._p.append(h)
d=Document(); s=d.sections[0]; s.page_width=Mm(210); s.page_height=Mm(297)
for a in("left_margin","right_margin"): setattr(s,a,Mm(17))
s.top_margin=Mm(14); s.bottom_margin=Mm(13)
st=d.styles["Normal"]; st.font.name="Arial"; st.font.size=Pt(9.5); st.paragraph_format.space_after=Pt(1); st.paragraph_format.space_before=Pt(0)
def para(text="",bold=False,size=None,color=None,after=None):
    p=d.add_paragraph(); r=p.add_run(text); r.bold=bold
    if size: r.font.size=Pt(size)
    if color: r.font.color.rgb=RGBColor.from_string(color)
    if after is not None: p.paragraph_format.space_after=Pt(after)
    return p
def heading(t):
    p=para(t.upper(),True,10.5,"0B1630"); p.paragraph_format.space_before=Pt(8)
    pPr=p._p.get_or_add_pPr(); b=OxmlElement("w:pBdr"); bt=OxmlElement("w:bottom")
    for k,v in (("val","single"),("sz","6"),("space","1"),("color","2F6BFF")): bt.set(qn("w:"+k),v)
    b.append(bt); pPr.append(b)
def bullet(t):
    p=d.add_paragraph(style="List Bullet"); p.add_run(t); p.paragraph_format.space_after=Pt(1)
para(P["name"],True,20,"0B1630",0); para(P["title"],False,10.5,"2F6BFF",2)
p=d.add_paragraph(); hyper(p,"mailto:"+P["email"],P["email"]); p.add_run("  |  ")
hyper(p,P["github"],"github.com/raunaksing15-blip"); p.add_run("  |  ")
hyper(p,P["linkedin"],"linkedin.com/in/raunak-singh-923298387"); p.add_run("  |  "); hyper(p,P["x"],"x.com/raunaksingh")
heading("Summary"); para(P["summary"])
e=P["education"]; heading("Education"); p=para(); p.add_run(e["school"]).bold=True; p.add_run(f" - {e['dates']}")
para(f"{e['degree']}. {e['note']}")
heading("Projects")
for pr in P["projects"]:
    p=para(); p.add_run(pr["name"]).bold=True; p.add_run(f" - {pr['tag']} | {pr['dates']}")
    p=d.add_paragraph()
    for i,(t,u) in enumerate(pr["links"]):
        if i: p.add_run("  |  ")
        hyper(p,u,t)
    for b in pr["bullets"]: bullet(b)
heading("Leadership Experience")
for l in P["leadership"]:
    p=para(); p.add_run(l["role"]).bold=True; p.add_run(f", {l['name']} ({l['place']}) | {l['dates']}")
    for b in l["bullets"]: bullet(b)
heading("Technical Skills")
for k,v in P["skills"]:
    p=d.add_paragraph(); p.add_run(k+": ").bold=True; p.add_run(v)
heading("Open Source and Interests"); para(P["interests"])
d.core_properties.author="Raunak Singh"; d.core_properties.title="Raunak Singh - Resume"
d.save(OUT+"Raunak_Singh_Resume.docx")
# plain text
e=P["education"]; o=[P["name"].upper(),P["title"],"",
 f"Email: {P['email']}",f"GitHub: {P['github']}",f"LinkedIn: {P['linkedin']}",f"X: {P['x']}","","SUMMARY",P["summary"],"",
 "EDUCATION",f"{e['school']} - {e['dates']}",f"{e['degree']}. {e['note']}","","PROJECTS"]
for pr in P["projects"]:
    o+= [f"{pr['name']} - {pr['tag']} | {pr['dates']}"]+[f"  {t}: {u}" for t,u in pr["links"]]+[f"  - {b}" for b in pr["bullets"]]+[""]
o+=["LEADERSHIP EXPERIENCE"]
for l in P["leadership"]: o+=[f"{l['role']}, {l['name']} ({l['place']}) | {l['dates']}"]+[f"  - {b}" for b in l["bullets"]]
o+=["","TECHNICAL SKILLS"]+[f"{k}: {v}" for k,v in P["skills"]]+["","OPEN SOURCE AND INTERESTS",P["interests"],""]
open(OUT+"Raunak_Singh_Resume.txt","w").write("\n".join(o))
